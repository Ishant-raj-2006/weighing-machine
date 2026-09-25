import io
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x14, 0x21, 0x3D)
STEEL = RGBColor(0x4A, 0x55, 0x68)
PASS_GREEN = RGBColor(0x1F, 0x7A, 0x4D)
FAIL_RED = RGBColor(0xB2, 0x3A, 0x34)


def _shade_cell(cell, hex_color):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_text(cell, text, bold=False, color=None, size=9, align_center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    if align_center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(str(text))
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def _add_kv_table(doc, pairs):
    table = doc.add_table(rows=0, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i in range(0, len(pairs), 2):
        row = table.add_row()
        left = pairs[i]
        right = pairs[i + 1] if i + 1 < len(pairs) else ("", "")
        _set_cell_text(row.cells[0], left[0], bold=True, color=STEEL)
        _set_cell_text(row.cells[1], left[1])
        _set_cell_text(row.cells[2], right[0], bold=True, color=STEEL)
        _set_cell_text(row.cells[3], right[1])
    return table


def _add_data_table(doc, headers, rows, result_col_index=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        _set_cell_text(table.rows[0].cells[i], h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align_center=True)
        _shade_cell(table.rows[0].cells[i], "14213D")
    for r in rows:
        row = table.add_row()
        for i, val in enumerate(r):
            color = None
            bold = False
            if result_col_index is not None and i == result_col_index:
                bold = True
                color = PASS_GREEN if val == "PASS" else (FAIL_RED if val == "FAIL" else STEEL)
            _set_cell_text(row.cells[i], val, bold=bold, color=color, align_center=True)
    return table


def build_docx(context: dict) -> bytes:
    report = context["report"]
    instrument = context["instrument"]

    doc = Document()
    section = doc.sections[0]
    section.left_margin = section.right_margin = Cm(1.6)
    section.top_margin = section.bottom_margin = Cm(1.6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("TEST REPORT FOR NON-AUTOMATIC WEIGHING INSTRUMENT (NAWI)")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = INK

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run("Issued in accordance with OIML Recommendation R 76-1  |  Legal Metrology Act, 2009")
    srun.font.size = Pt(9.5)
    srun.font.color.rgb = STEEL

    def heading(text):
        h = doc.add_heading(level=2)
        r = h.add_run(text)
        r.font.color.rgb = INK
        r.font.size = Pt(12)

    heading("1. Report & Laboratory Details")
    _add_kv_table(doc, [
        ("Report No.", report["report_number"]),
        ("Test Date", str(report["test_date"])),
        ("Laboratory", report["lab_name"]),
        ("Test Stage", report["test_stage"].replace("_", " ").title()),
        ("Ambient Temperature", f'{report.get("lab_temperature_c", "-")} °C'),
        ("Relative Humidity", f'{report.get("lab_humidity_pct", "-")} %'),
        ("Atmospheric Pressure", f'{report.get("atmospheric_pressure_hpa", "-")} hPa'),
        ("Status", report["status"].replace("_", " ").title()),
    ])

    heading("2. Instrument Under Test")
    _add_kv_table(doc, [
        ("Manufacturer", instrument["manufacturer_name"]),
        ("Model", instrument["model_name"]),
        ("Serial No.", instrument["serial_number"]),
        ("Instrument Type", instrument["instrument_type"]),
        ("Max Capacity (Max)", f'{instrument["max_capacity"]} {instrument["unit"]}'),
        ("Min Capacity (Min)", f'{instrument["min_capacity"]} {instrument["unit"]}'),
        ("Verification Scale Interval (e)", f'{instrument["e_value"]} {instrument["unit"]}'),
        ("Accuracy Class", f'Class {instrument["accuracy_class"]}'),
    ])

    weighing = context.get("weighing", [])
    if weighing:
        heading("3. Accuracy (Weighing) Test")
        unit = instrument["unit"]
        rows = [[
            str(i + 1), f'{w["test_load"]} {unit}', f'{w["indicated_value"]} {unit}',
            f'{w["error"]:+.3f} {unit}', f'± {w["mpe"]:.3f} {unit}', f'{w["verification_intervals"]}',
            w["result"],
        ] for i, w in enumerate(weighing)]
        _add_data_table(doc, ["#", "Test Load", "Indication", "Error", "MPE", "m = load/e", "Result"],
                         rows, result_col_index=6)

    repeatability = context.get("repeatability", [])
    if repeatability:
        heading("4. Repeatability Test")
        unit = instrument["unit"]
        rows = [[
            f'{rt["test_load"]} {unit}', ", ".join(f'{v:g}' for v in rt["readings"]),
            f'{rt["mean_value"]:.3f} {unit}', f'{rt["range_value"]:.3f} {unit}',
            f'± {rt["mpe"]:.3f} {unit}', rt["result"],
        ] for rt in repeatability]
        _add_data_table(doc, ["Test Load", "Readings", "Mean", "Range", "MPE", "Result"],
                         rows, result_col_index=5)

    eccentricity = context.get("eccentricity", [])
    if eccentricity:
        heading("5. Eccentricity (Corner Load) Test")
        unit = instrument["unit"]
        rows = [[
            e["position"], f'{e["test_load"]} {unit}', f'{e["indicated_value"]} {unit}',
            f'{e["error"]:+.3f} {unit}', f'± {e["mpe"]:.3f} {unit}', e["result"],
        ] for e in eccentricity]
        _add_data_table(doc, ["Position", "Test Load", "Indication", "Error", "MPE", "Result"],
                         rows, result_col_index=5)

    doc.add_paragraph()
    overall = report.get("overall_result") or "PENDING"
    color = PASS_GREEN if overall == "PASS" else (FAIL_RED if overall == "FAIL" else STEEL)
    banner_table = doc.add_table(rows=1, cols=1)
    cell = banner_table.rows[0].cells[0]
    _shade_cell(cell, "1F7A4D" if overall == "PASS" else ("B23A34" if overall == "FAIL" else "4A5568"))
    _set_cell_text(cell, f"OVERALL RESULT: {overall}", bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), size=13, align_center=True)

    if report.get("remarks"):
        heading("Remarks")
        doc.add_paragraph(report["remarks"])

    doc.add_paragraph()
    sign_table = doc.add_table(rows=3, cols=2)
    _set_cell_text(sign_table.rows[0].cells[0], "Tested By", bold=True, color=STEEL)
    _set_cell_text(sign_table.rows[0].cells[1], "Reviewed By", bold=True, color=STEEL)
    _set_cell_text(sign_table.rows[1].cells[0], context.get("tested_by_name") or "-")
    _set_cell_text(sign_table.rows[1].cells[1], context.get("reviewed_by_name") or "-")
    _set_cell_text(sign_table.rows[2].cells[0], "Signature: ______________________")
    _set_cell_text(sign_table.rows[2].cells[1], "Signature: ______________________")

    footer = doc.add_paragraph()
    frun = footer.add_run(
        "This report was generated by the NAWI Test Report system. Prototype scope — verify MPE "
        "values against the current official OIML R76-1 text before use beyond demonstration."
    )
    frun.font.size = Pt(8)
    frun.font.color.rgb = STEEL

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
